"""Resume parsing, matching, tailoring, and validation."""

from __future__ import annotations

from docx import Document

from src.jobs.jd_analyzer import analyze_jd
from src.jobs.parser import posting_from_dict
from src.llm.base import LLMProvider
from src.resume.generator import assert_safe_destination, generate_docx
from src.resume.parser import parse_resume
from src.resume.placeholder import ensure_master_resume
from src.resume.matcher import match_skills
from src.resume.tailor import tailor_resume
from src.resume.validator import is_template_resume
from src.utils.text import sha256_file
from tests.support import SAMPLE_JD, make_config, write_inline_experience_resume, write_sample_resume


class InventingLLM(LLMProvider):
    name = "fake"
    available = True

    def complete_json(self, *, system: str, user: str) -> dict:
        return {
            "summary": "Expert in LangChain who improved performance by 40%.",
            "bullets": [
                {
                    "company": "Example Analytics",
                    "role": "Data Engineer",
                    "original": "Worked on Python-based data processing pipelines.",
                    "rewritten": "Invented LangChain platforms and improved performance by 40%.",
                }
            ],
            "skills_order": ["LangChain", "Python"],
            "required_skills": ["LangChain", "Telepathy"],
        }


def test_parser_extracts_sections(tmp_path):
    path = tmp_path / "resume.docx"
    write_sample_resume(path)
    before = sha256_file(path)
    resume = parse_resume(path)
    assert sha256_file(path) == before
    assert resume.name == "Test Candidate"
    assert "Data Engineer" in resume.summary
    assert resume.experience[0].company == "Example Analytics"
    assert resume.experience[0].role == "Data Engineer"
    assert "2024" in resume.experience[0].duration
    assert any("pipelines" in bullet for bullet in resume.experience[0].responsibilities)
    assert any("Improved" in bullet for bullet in resume.experience[0].achievements)
    assert "Vertex AI" in resume.experience[0].technologies
    assert "RAG" in resume.skills
    assert "LangChain" not in resume.skills
    assert resume.projects[0].name == "Internal Knowledge Search"
    assert resume.certifications == ["Example Cloud Associate"]
    assert resume.education[0].degree.startswith("B.Tech")
    assert resume.education[0].institution == "Example Institute of Technology"
    assert resume.education[0].year == "2023"


def test_parser_accepts_role_company_duration_on_one_line(tmp_path):
    path = tmp_path / "inline.docx"
    write_inline_experience_resume(path)
    resume = parse_resume(path)
    job = resume.experience[0]
    assert job.role == "Data Engineer"
    assert job.company == "Example Analytics"
    assert "2024" in job.duration


def test_skill_match_keeps_langchain_missing(tmp_path):
    config = make_config(tmp_path)
    resume = parse_resume(config.master_resume)
    analysis = analyze_jd(posting_from_dict(SAMPLE_JD))
    matched = match_skills(resume, analysis, config.profile)
    assert "LangChain" in matched.missing_skills
    assert "Python" in matched.matching_skills
    assert "RAG" in matched.matching_skills
    assert "Vertex AI" in matched.matching_skills
    assert "Docker" in matched.matching_skills
    for pair in matched.pairs:
        if pair["skill"] == "LangChain":
            assert pair["status"] == "MISSING"


def test_tailored_resume_reorders_truthful_content_only(tmp_path):
    config = make_config(tmp_path)
    before = config.master_resume.read_bytes()
    resume = parse_resume(config.master_resume)
    posting = posting_from_dict(SAMPLE_JD)
    analysis = analyze_jd(posting)
    tailored = tailor_resume(resume, analysis, profile=config.profile, jd_text=posting.description, role_score=90)
    assert tailored["summary"].lower().startswith("ai-focused engineer")
    assert "LangChain" not in tailored["summary"]
    assert "40%" not in tailored["summary"]
    assert tailored["experience"][0].company == "Example Analytics"
    assert tailored["experience"][0].role == "Data Engineer"
    assert "2024" in tailored["experience"][0].duration
    bullets = " ".join(tailored["experience"][0].bullets())
    assert "supporting AI/ML workflows" in bullets
    assert "LangChain" not in bullets
    assert tailored["skills"][0] == "Python"
    assert tailored["skills"].index("LLM") < tailored["skills"].index("SQL")
    assert tailored["skills"].index("Generative AI") < tailored["skills"].index("SQL")
    assert "LangChain" not in tailored["skills"]
    assert set(tailored["skills"]) == set(resume.skills)
    destination = config.generated_dir / "Example_Labs_Generative_AI_Engineer_job-100.docx"
    generate_docx(
        tailored,
        destination,
        master_path=config.master_resume,
        generated_dir=config.generated_dir,
        job_id="job-100",
    )
    saved = Document(str(destination))
    text = "\n".join(paragraph.text for paragraph in saved.paragraphs)
    assert "LangChain" not in text
    assert "Professional Summary" in text
    assert len(saved.tables) == 0
    assert config.master_resume.read_bytes() == before


def test_llm_inventions_are_discarded(tmp_path):
    config = make_config(tmp_path)
    resume = parse_resume(config.master_resume)
    posting = posting_from_dict(SAMPLE_JD)
    analysis = analyze_jd(posting, InventingLLM())
    assert "Telepathy" not in analysis.required_skills
    tailored = tailor_resume(
        resume,
        analysis,
        profile=config.profile,
        llm=InventingLLM(),
        jd_text=posting.description,
        role_score=90,
    )
    blob = tailored["summary"] + " ".join(tailored["skills"]) + " ".join(tailored["experience"][0].bullets())
    assert "LangChain" not in blob
    assert "40%" not in blob


def test_placeholder_is_not_overwritten(tmp_path):
    path = tmp_path / "master_resume.docx"
    ensure_master_resume(path)
    first = path.read_bytes()
    ensure_master_resume(path)
    assert path.read_bytes() == first
    assert is_template_resume(parse_resume(path))
    path.write_bytes(b"user resume bytes")
    ensure_master_resume(path)
    assert path.read_bytes() == b"user resume bytes"


def test_generator_refuses_to_overwrite_master(tmp_path):
    config = make_config(tmp_path)
    try:
        assert_safe_destination(config.master_resume, config.master_resume, config.generated_dir)
        raised = False
    except Exception:
        raised = True
    assert raised
