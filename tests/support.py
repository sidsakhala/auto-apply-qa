"""Shared fixtures for resume and configuration tests."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from src.utils.config import PROJECT_ROOT, load_config


def write_sample_resume(path: Path) -> None:
    document = Document()
    document.add_heading("Test Candidate", 0)
    document.add_heading("Professional Summary", level=1)
    document.add_paragraph(
        "Data Engineer with experience in Python, SQL, cloud technologies, and building data processing pipelines. "
        "Hands-on with Docker, Kubernetes, Vertex AI, RAG workflows, LLM applications, and machine learning."
    )
    document.add_heading("Work Experience", level=1)
    role = document.add_paragraph("Data Engineer")
    role.runs[0].bold = True
    document.add_paragraph("Example Analytics | Jan 2024 – Present")
    document.add_paragraph("Worked on Python-based data processing pipelines.", style="List Bullet")
    document.add_paragraph("Built RAG workflows on Vertex AI for internal search.", style="List Bullet")
    document.add_paragraph("Containerized services with Docker and deployed them on Kubernetes.", style="List Bullet")
    document.add_paragraph("Improved batch refresh time for internal datasets.", style="List Bullet")
    document.add_paragraph("Technologies: Python, SQL, Vertex AI, Docker, Kubernetes, RAG, GCP")
    document.add_heading("Projects", level=1)
    project = document.add_paragraph("Internal Knowledge Search")
    project.runs[0].bold = True
    document.add_paragraph("Implemented retrieval workflows over internal documents.", style="List Bullet")
    document.add_heading("Skills", level=1)
    document.add_paragraph(
        "Python, SQL, PySpark, Kafka, AWS, Azure, GCP, Vertex AI, Docker, Kubernetes, Terraform, "
        "Machine Learning, Generative AI, LLM, RAG, MLflow"
    )
    document.add_heading("Certifications", level=1)
    document.add_paragraph("Example Cloud Associate", style="List Bullet")
    document.add_heading("Education", level=1)
    degree = document.add_paragraph("B.Tech, Computer Engineering")
    degree.runs[0].bold = True
    document.add_paragraph("Example Institute of Technology | 2023")
    document.save(str(path))


def write_inline_experience_resume(path: Path) -> None:
    document = Document()
    document.add_heading("Test Candidate", 0)
    document.add_heading("Work Experience", level=1)
    document.add_paragraph("Data Engineer | Example Analytics | Jan 2024 – Present")
    document.add_paragraph("Built internal Python tools.", style="List Bullet")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python")
    document.save(str(path))


def make_config(tmp_path: Path):
    config = load_config(PROJECT_ROOT)
    config.master_resume = tmp_path / "master_resume.docx"
    config.generated_dir = tmp_path / "generated"
    config.database = tmp_path / "jobs.db"
    config.reports = tmp_path / "reports"
    config.logs = tmp_path / "logs"
    config.rate_limits.min_delay_seconds = 0
    config.rate_limits.max_delay_seconds = 0
    config.overleaf_project_url = ""
    config.overleaf_dir = tmp_path / "overleaf"
    config.ensure_dirs()
    write_sample_resume(config.master_resume)
    return config


SAMPLE_JD = {
    "job_id": "job-100",
    "job_title": "QA Automation Engineer",
    "company": "Example Labs",
    "location": "Pune",
    "salary": "₹12–18 LPA",
    "experience": "2-5 years",
    "job_url": "https://www.naukri.com/qa-automation-engineer-example-100000001",
    "description": """
Example Labs is hiring a QA Automation Engineer.

Responsibilities:
- Build Selenium and TestNG automation frameworks.
- API testing with Postman and Rest Assured.
- CI/CD with Jenkins.

Required skills:
- Selenium
- Java
- TestNG
- API Testing
- Postman

Preferred skills:
- Cypress
- Playwright
""",
}
