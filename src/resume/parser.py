"""Read ``resume/master_resume.docx`` into structured sections.

The parser only reads. It never saves the document it opened.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from src.exceptions import ConfigurationError
from src.resume.models import Education, Experience, Project, Resume
from src.utils.text import normalize_space, unique_keep

SECTION_ALIASES = {
    "summary": {"professional summary", "summary", "profile", "about me", "about"},
    "experience": {
        "work experience",
        "professional experience",
        "experience",
        "employment history",
        "employment",
    },
    "projects": {"projects", "personal projects", "selected projects", "academic projects"},
    "skills": {"skills", "technical skills", "core skills", "skills and technologies"},
    "certifications": {"certifications", "certificates", "licenses", "licenses and certifications"},
    "education": {"education", "academic background", "academics"},
}

_ACHIEVEMENT = re.compile(
    r"\b(achieved|awarded|improved|reduced|increased|decreased|saved|won|recognized)\b",
    re.IGNORECASE,
)
_TECH_LINE = re.compile(r"^(technologies|tech stack|tools)\s*:\s*(.+)$", re.IGNORECASE)
_BULLET_PREFIX = re.compile(r"^[\u2022\-\*\u25aa]\s+")
_ROLE_HINT = re.compile(
    r"\b(engineer|developer|analyst|scientist|manager|intern|consultant|lead|architect|designer)\b",
    re.IGNORECASE,
)


def _style_name(paragraph: Paragraph) -> str:
    style = paragraph.style
    return style.name if style is not None else ""


def _is_bold(paragraph: Paragraph) -> bool:
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(bool(run.bold) for run in runs)


def _is_list(paragraph: Paragraph) -> bool:
    style = _style_name(paragraph).lower()
    if "list" in style:
        return True
    return paragraph._p.find(qn("w:numPr")) is not None


def _section_for(text: str) -> str | None:
    key = normalize_space(text).lower().rstrip(":")
    for name, aliases in SECTION_ALIASES.items():
        if key in aliases:
            return name
    return None


def _looks_like_duration(text: str) -> bool:
    lowered = text.lower()
    if re.search(r"(19|20)\d{2}", lowered):
        return True
    return bool(
        re.search(
            r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec|present|current)\b",
            lowered,
        )
    )


def _split_meta(text: str) -> list[str]:
    if "|" in text:
        parts = re.split(r"\s*\|\s*", text)
    else:
        parts = re.split(r"\s+[-–—]\s+", text)
    return [normalize_space(part) for part in parts if normalize_space(part)]


def _split_list(text: str) -> list[str]:
    chunks = re.split(r"[,;|•\n]+", text)
    return unique_keep([normalize_space(chunk) for chunk in chunks if normalize_space(chunk)])


def _clean_bullet(text: str) -> str:
    return normalize_space(_BULLET_PREFIX.sub("", text))


def _is_bullet(paragraph: Paragraph, text: str) -> bool:
    return _is_list(paragraph) or bool(_BULLET_PREFIX.match(text))


class ResumeParser:
    """Parse a DOCX resume that uses standard headings and bullet paragraphs."""

    def parse(self, path: Path | str) -> Resume:
        path = Path(path)
        if not path.exists():
            raise ConfigurationError(
                f"Master resume not found at {path}. Place your resume there and run again."
            )
        document = Document(str(path))
        paragraphs = [item for item in document.paragraphs if normalize_space(item.text)]

        name = ""
        section = ""
        summary: list[str] = []
        skills: list[str] = []
        certifications: list[str] = []
        experience: list[Experience] = []
        projects: list[Project] = []
        education: list[Education] = []
        current_exp: Experience | None = None
        current_project: Project | None = None
        current_edu: Education | None = None

        def flush_education() -> None:
            nonlocal current_edu
            if current_edu and current_edu.degree:
                education.append(current_edu)
            current_edu = None

        for index, paragraph in enumerate(paragraphs):
            text = normalize_space(paragraph.text)
            style = _style_name(paragraph)
            detected = _section_for(text)
            heading = style.startswith("Heading") or style == "Title"
            if detected and (heading or _is_bold(paragraph) or len(text) <= 48):
                flush_education()
                section = detected
                current_exp = None
                current_project = None
                continue

            if index == 0 and not name and not detected:
                name = text
                continue

            if section == "summary":
                summary.append(text)
            elif section == "skills":
                skills.extend(_split_list(_clean_bullet(text)))
            elif section == "certifications":
                certifications.append(_clean_bullet(text))
            elif section == "experience":
                current_exp = _consume_experience(paragraph, text, current_exp, experience)
            elif section == "projects":
                current_project = _consume_project(paragraph, text, current_project, projects)
            elif section == "education":
                current_edu = _consume_education(paragraph, text, current_edu, education)

        flush_education()
        experience = [item for item in experience if item.role or item.company or item.bullets()]
        projects = [item for item in projects if item.name or item.bullets]
        raw_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return Resume(
            name=name,
            summary=" ".join(summary).strip(),
            experience=experience,
            projects=projects,
            skills=unique_keep(skills),
            certifications=unique_keep([item for item in certifications if item]),
            education=education,
            raw_text=raw_text,
        )


def _consume_experience(
    paragraph: Paragraph,
    text: str,
    current: Experience | None,
    experience: list[Experience],
) -> Experience | None:
    tech = _TECH_LINE.match(text)
    if tech and current is not None:
        current.technologies.extend(_split_list(tech.group(2)))
        return current

    if _is_bullet(paragraph, text):
        if current is None:
            current = Experience(company="", role="", duration="")
            experience.append(current)
        bullet = _clean_bullet(text)
        if _ACHIEVEMENT.search(bullet):
            current.achievements.append(bullet)
        else:
            current.responsibilities.append(bullet)
        return current

    parts = _split_meta(text)
    durations = [part for part in parts if _looks_like_duration(part)]
    if durations and len(parts) >= 2:
        duration = durations[-1]
        rest = [part for part in parts if part != duration]
        open_role = (
            current is not None
            and current.role
            and not current.company
            and not current.bullets()
        )
        if open_role and len(rest) == 1:
            current.company = rest[0]
            current.duration = duration
            return current
        company, role = _company_and_role(rest)
        entry = Experience(company=company, role=role, duration=duration)
        experience.append(entry)
        return entry

    if current is not None and current.role and not current.company and not current.bullets():
        current.company = text
        return current

    entry = Experience(company="", role=text, duration="")
    experience.append(entry)
    return entry


def _company_and_role(rest: list[str]) -> tuple[str, str]:
    if len(rest) >= 2:
        first, second = rest[0], rest[1]
        if _ROLE_HINT.search(first) and not _ROLE_HINT.search(second):
            return second, first
        return first, second
    if len(rest) == 1:
        return rest[0], ""
    return "", ""


def _consume_project(
    paragraph: Paragraph,
    text: str,
    current: Project | None,
    projects: list[Project],
) -> Project | None:
    tech = _TECH_LINE.match(text)
    if tech and current is not None:
        current.technologies.extend(_split_list(tech.group(2)))
        return current
    if _is_bullet(paragraph, text):
        if current is None:
            current = Project(name="")
            projects.append(current)
        current.bullets.append(_clean_bullet(text))
        return current
    current = Project(name=text)
    projects.append(current)
    return current


def _consume_education(
    paragraph: Paragraph,
    text: str,
    current: Education | None,
    education: list[Education],
) -> Education | None:
    if _is_bold(paragraph) or current is None:
        if current is not None and current.degree:
            education.append(current)
        return Education(degree=_clean_bullet(text))

    if not current.institution:
        parts = _split_meta(text)
        if len(parts) >= 2 and _looks_like_duration(parts[-1]):
            current.institution = parts[0]
            current.year = parts[-1]
        else:
            current.institution = text
        return current

    current.details = normalize_space(f"{current.details} {text}")
    return current


def parse_resume(path: Path | str) -> Resume:
    return ResumeParser().parse(path)
