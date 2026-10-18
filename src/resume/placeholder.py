"""Create a placeholder master resume when the user has not added one yet.

An existing master resume is never replaced.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from src.resume.validator import TEMPLATE_MARKER


def write_placeholder(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("YOUR NAME", 0)
    document.add_heading("Professional Summary", level=1)
    document.add_paragraph(
        "Replace this summary with your real professional summary. "
        f"This placeholder must not be used to apply. {TEMPLATE_MARKER}"
    )
    document.add_heading("Work Experience", level=1)
    role = document.add_paragraph("Your Role")
    role.runs[0].bold = True
    document.add_paragraph("Your Company | Jan 2024 – Present")
    document.add_paragraph("Describe a real responsibility from your work.", style="List Bullet")
    document.add_paragraph("Technologies: Python")
    document.add_heading("Projects", level=1)
    project = document.add_paragraph("Your Project")
    project.runs[0].bold = True
    document.add_paragraph("Describe a real project you worked on.", style="List Bullet")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python")
    document.add_heading("Certifications", level=1)
    document.add_paragraph("Your certification", style="List Bullet")
    document.add_heading("Education", level=1)
    degree = document.add_paragraph("Your Degree")
    degree.runs[0].bold = True
    document.add_paragraph("Your Institution | 2024")
    document.save(str(path))


def ensure_master_resume(path: Path) -> None:
    """Create the placeholder only when the master file is absent."""
    if path.exists():
        return
    write_placeholder(path)
