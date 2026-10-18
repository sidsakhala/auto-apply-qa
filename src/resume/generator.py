"""Write an ATS-friendly DOCX. This module never opens the master resume for writing."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from src.exceptions import ResumeIntegrityError
from src.resume.models import Education, Experience, Project

SECTIONS = (
    "Professional Summary",
    "Skills",
    "Work Experience",
    "Projects",
    "Education",
    "Certifications",
)


def assert_safe_destination(destination: Path, master_path: Path, generated_dir: Path) -> Path:
    destination = destination.resolve()
    master_path = master_path.resolve()
    generated_dir = generated_dir.resolve()
    if destination == master_path:
        raise ResumeIntegrityError("Refusing to overwrite the master resume")
    if generated_dir not in destination.parents:
        raise ResumeIntegrityError(f"Refusing to write a resume outside {generated_dir}")
    return destination


def _set_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for heading_name in ("Heading 1", "Title"):
        heading = document.styles[heading_name]
        heading.font.name = "Calibri"


def _add_bullets(document: Document, bullets: list[str]) -> None:
    for bullet in bullets:
        if bullet.strip():
            document.add_paragraph(bullet.strip(), style="List Bullet")


def generate_docx(
    tailored: dict,
    destination: Path,
    *,
    master_path: Path,
    generated_dir: Path,
    job_id: str,
) -> Path:
    destination = assert_safe_destination(Path(destination), Path(master_path), Path(generated_dir))
    destination.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _set_font(document)
    for section in document.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    document.add_heading(tailored.get("name") or "Resume", 0)
    document.add_heading("Professional Summary", level=1)
    document.add_paragraph(tailored["summary"])

    document.add_heading("Skills", level=1)
    document.add_paragraph(", ".join(tailored["skills"]))

    document.add_heading("Work Experience", level=1)
    for item in tailored["experience"]:
        assert isinstance(item, Experience)
        title = document.add_paragraph(item.role or item.company)
        if title.runs:
            title.runs[0].bold = True
        meta = " | ".join(part for part in (item.company, item.duration) if part)
        if meta and meta != item.role:
            document.add_paragraph(meta)
        _add_bullets(document, item.bullets())
        if item.technologies:
            document.add_paragraph("Technologies: " + ", ".join(item.technologies))

    if tailored["projects"]:
        document.add_heading("Projects", level=1)
        for project in tailored["projects"]:
            assert isinstance(project, Project)
            name = document.add_paragraph(project.name)
            if name.runs:
                name.runs[0].bold = True
            _add_bullets(document, project.bullets)
            if project.technologies:
                document.add_paragraph("Technologies: " + ", ".join(project.technologies))

    if tailored["education"]:
        document.add_heading("Education", level=1)
        for item in tailored["education"]:
            assert isinstance(item, Education)
            degree = document.add_paragraph(item.degree)
            if degree.runs:
                degree.runs[0].bold = True
            meta = " | ".join(part for part in (item.institution, item.year) if part)
            if meta:
                document.add_paragraph(meta)
            if item.details:
                document.add_paragraph(item.details)

    if tailored["certifications"]:
        document.add_heading("Certifications", level=1)
        _add_bullets(document, tailored["certifications"])

    document.core_properties.subject = f"tailored-resume:{job_id}"[:200]
    document.core_properties.category = "tailored-resume"
    document.core_properties.comments = "Generated from the master resume. Facts were not invented."
    document.save(str(destination))
    return destination
